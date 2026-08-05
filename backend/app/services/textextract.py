"""Stage 1 — text extraction with page/block geometry.

PDF: pdfplumber word geometry + visual Arabic de-shaping; PyMuPDF fallback if empty.
DOCX: python-docx with paragraph direction; pseudo-pages for markers.
"""
from __future__ import annotations

import io
import logging
import re
from zipfile import BadZipFile
from typing import Any

from .bidi_text import deshape_word, detect_direction

logger = logging.getLogger(__name__)


class TextExtractError(Exception):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


def _pdf_fitz_plain(data: bytes) -> list[dict]:
    import fitz

    pages = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            pages.append({"page": i, "text": page.get_text() or ""})
    return pages


def _cluster_lines(words: list[dict], y_tol: float = 3.0) -> list[list[dict]]:
    if not words:
        return []
    sorted_w = sorted(words, key=lambda w: (w["top"], w["x0"]))
    lines: list[list[dict]] = []
    current: list[dict] = [sorted_w[0]]
    ref_top = sorted_w[0]["top"]
    for w in sorted_w[1:]:
        if abs(w["top"] - ref_top) <= y_tol:
            current.append(w)
        else:
            lines.append(current)
            current = [w]
            ref_top = w["top"]
    lines.append(current)
    return lines


def _cluster_columns(lines: list[list[dict]], page_width: float, gap_ratio: float = 0.08) -> list[list[list[dict]]]:
    if not lines:
        return []
    if len(lines) == 1:
        return [lines]
    centers = [sum(w["x0"] + w["x1"] for w in ln) / (2 * len(ln)) for ln in lines]
    gap_thresh = page_width * gap_ratio
    columns: list[list[list[dict]]] = [[lines[0]]]
    col_centers = [centers[0]]
    for i in range(1, len(lines)):
        c = centers[i]
        best_j = min(range(len(col_centers)), key=lambda j: abs(c - col_centers[j]))
        if abs(c - col_centers[best_j]) <= gap_thresh:
            columns[best_j].append(lines[i])
            col_centers[best_j] = (col_centers[best_j] + c) / 2
        else:
            columns.append([lines[i]])
            col_centers.append(c)
    return columns


def _line_bbox(line: list[dict]) -> list[float]:
    x0 = min(w["x0"] for w in line)
    x1 = max(w["x1"] for w in line)
    top = min(w["top"] for w in line)
    bottom = max(w["bottom"] for w in line)
    return [x0, top, x1, bottom]


def _words_to_text(words: list[dict], direction: str) -> str:
    if direction == "rtl":
        ordered = sorted(words, key=lambda w: -w["x0"])
    else:
        ordered = sorted(words, key=lambda w: w["x0"])
    return " ".join(w["text"] for w in ordered)


def _block_from_lines(lines: list[list[dict]], column: int) -> dict[str, Any]:
    all_words = [w for ln in lines for w in ln]
    text_parts: list[str] = []
    for ln in lines:
        line_text = _words_to_text(ln, detect_direction(" ".join(w["text"] for w in ln)))
        text_parts.append(line_text)
    text = "\n".join(text_parts)
    direction = detect_direction(text)
    bbox = [
        min(w["x0"] for w in all_words),
        min(w["top"] for w in all_words),
        max(w["x1"] for w in all_words),
        max(w["bottom"] for w in all_words),
    ]
    return {"text": text, "bbox": bbox, "direction": direction, "column": column}


def _split_line_by_gaps(line: list[dict], page_width: float, gap_ratio: float = 0.06) -> list[list[dict]]:
    if len(line) <= 1:
        return [line]
    ordered = sorted(line, key=lambda w: w["x0"])
    groups: list[list[dict]] = [[ordered[0]]]
    for w in ordered[1:]:
        prev = groups[-1][-1]
        if w["x0"] - prev["x1"] > page_width * gap_ratio:
            groups.append([w])
        else:
            groups[-1].append(w)
    return groups


def _pdf_blocks(data: bytes) -> list[dict]:
    import pdfplumber

    layout_pages: list[dict] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            raw_words = page.extract_words() or []
            words = []
            for w in raw_words:
                t = (w.get("text") or "").strip()
                if not t:
                    continue
                words.append(
                    {
                        "text": deshape_word(t),
                        "x0": float(w["x0"]),
                        "x1": float(w["x1"]),
                        "top": float(w["top"]),
                        "bottom": float(w["bottom"]),
                    }
                )
            pw = float(page.width)
            ph = float(page.height)
            line_groups = _cluster_lines(words)
            raw_blocks: list[dict] = []
            for line in line_groups:
                for group in _split_line_by_gaps(line, pw):
                    if group:
                        raw_blocks.append(_block_from_lines([group], 1))
            page_dir = detect_direction(" ".join(w["text"] for w in words))
            raw_blocks.sort(
                key=lambda b: b["bbox"][0] + b["bbox"][2],
                reverse=(page_dir in ("rtl", "mixed")),
            )
            for i, b in enumerate(raw_blocks, start=1):
                b["column"] = i
            blocks = raw_blocks
            text = "\n".join(b["text"] for b in blocks)
            layout_pages.append(
                {"page": i, "width": pw, "height": ph, "blocks": blocks, "text": text}
            )
    return layout_pages


def _docx_blocks(data: bytes) -> list[dict]:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.oxml.ns import qn

    if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise TextExtractError("encrypted_or_unsupported_docx")

    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, BadZipFile, KeyError):
        raise TextExtractError("corrupted")

    blocks: list[dict] = []

    def _para_dir(para) -> str:
        p_pr = para._element.pPr
        if p_pr is not None:
            bidi = p_pr.find(qn("w:bidi"))
            if bidi is not None:
                value = bidi.get(qn("w:val"))
                if value is None or value.strip().lower() not in {"0", "false", "off", "no"}:
                    return "rtl"
        return detect_direction(para.text)

    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        direction = _para_dir(para)
        blocks.append(
            {
                "text": para.text,
                "bbox": [0.0, 0.0, 0.0, 0.0],
                "direction": direction,
                "column": 1,
            }
        )
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                t = " | ".join(cells)
                blocks.append(
                    {"text": t, "bbox": [0.0, 0.0, 0.0, 0.0], "direction": detect_direction(t), "column": 1}
                )

    if not blocks:
        if doc.paragraphs or doc.tables:
            raise TextExtractError("no_meaningful_text")
        raise TextExtractError("empty_document")

    pages: list[dict] = []
    current_blocks: list[dict] = []
    size = 0
    page_num = 1
    for b in blocks:
        current_blocks.append(b)
        size += len(b["text"]) + 1
        if size >= 2500:
            text = "\n".join(x["text"] for x in current_blocks)
            pages.append(
                {
                    "page": page_num,
                    "width": 595.0,
                    "height": 842.0,
                    "blocks": current_blocks,
                    "text": text,
                }
            )
            page_num += 1
            current_blocks, size = [], 0
    if current_blocks:
        text = "\n".join(x["text"] for x in current_blocks)
        pages.append(
            {
                "page": page_num,
                "width": 595.0,
                "height": 842.0,
                "blocks": current_blocks,
                "text": text,
            }
        )
    return pages


def extract_pages_geometry(data: bytes, filename: str) -> list[dict]:
    """Return [{page, width, height, blocks, text}]."""
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            pages = _pdf_blocks(data)
            joined = "".join(p.get("text", "") for p in pages)
            if not joined.strip():
                plain = _pdf_fitz_plain(data)
                joined = "".join(p["text"] for p in plain)
                if not joined.strip():
                    raise TextExtractError("scanned_pdf_not_supported")
                return [
                    {
                        "page": p["page"],
                        "width": 595.0,
                        "height": 842.0,
                        "blocks": [
                            {
                                "text": p["text"],
                                "bbox": [0.0, 0.0, 595.0, 842.0],
                                "direction": detect_direction(p["text"]),
                                "column": 1,
                            }
                        ],
                        "text": p["text"],
                    }
                    for p in plain
                ]
            return pages
        if name.endswith(".docx"):
            pages = _docx_blocks(data)
            return pages
    except TextExtractError:
        raise
    except Exception as exc:
        if name.endswith(".docx"):
            logger.exception(
                "docx_extraction_failed",
                extra={"document_type": "docx", "exception_type": type(exc).__name__},
            )
            raise TextExtractError("extraction_failed")
        raise TextExtractError("corrupted")
    raise TextExtractError("unsupported_type")


def extract_pages(data: bytes, filename: str) -> list[dict]:
    """Return [{page:int, text:str}] for pipeline compatibility."""
    geo = extract_pages_geometry(data, filename)
    return [{"page": p["page"], "text": p.get("text") or ""} for p in geo]


def build_raw_text(pages: list[dict]) -> tuple[str, dict[int, int]]:
    """Concatenate pages with [[PAGE n]] markers."""
    parts: list[str] = []
    starts: dict[int, int] = {}
    offset = 0
    for p in pages:
        marker = f"\n[[PAGE {p['page']}]]\n"
        starts[p["page"]] = offset
        parts.append(marker + p["text"])
        offset += len(marker) + len(p["text"])
    return "".join(parts), starts


def layout_from_geometry(geo_pages: list[dict]) -> list[dict]:
    """Strip to JSON-serializable layout stored on contracts.page_layout."""
    out = []
    for p in geo_pages:
        out.append(
            {
                "page": p["page"],
                "width": p.get("width"),
                "height": p.get("height"),
                "blocks": p.get("blocks") or [],
            }
        )
    return out
