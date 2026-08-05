import io
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient

from app.db import get_db
from app.main import app
from app.services.textextract import TextExtractError, _docx_blocks, extract_pages_geometry


def _docx_bytes(
    paragraphs: list[str],
    *,
    explicit_bidi: set[int] | None = None,
    add_paragraph_properties: bool = False,
) -> bytes:
    doc = Document()
    for index, text in enumerate(paragraphs):
        paragraph = doc.add_paragraph(text)
        if add_paragraph_properties:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if index in (explicit_bidi or set()):
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _blocks(data: bytes) -> list[dict]:
    return [block for page in _docx_blocks(data) for block in page["blocks"]]


def test_valid_english_docx_extracts_ltr():
    blocks = _blocks(_docx_bytes(["This agreement is effective today."], add_paragraph_properties=True))
    assert [(block["text"], block["direction"]) for block in blocks] == [
        ("This agreement is effective today.", "ltr")
    ]


def test_valid_arabic_docx_extracts_logical_rtl_text():
    text = "هذا العقد ساري المفعول"
    blocks = _blocks(_docx_bytes([text]))
    assert blocks[0]["text"] == text
    assert blocks[0]["direction"] == "rtl"


def test_mixed_arabic_english_docx_is_marked_mixed():
    blocks = _blocks(_docx_bytes(["عقد abc"]))
    assert blocks[0]["direction"] == "mixed"


def test_arabic_filename_extracts_normally():
    pages = extract_pages_geometry(_docx_bytes(["نص العقد"]), "عقد تجريبي.docx")
    assert pages[0]["text"] == "نص العقد"


def test_explicit_word_bidi_metadata_sets_rtl():
    blocks = _blocks(_docx_bytes(["English text with RTL metadata"], explicit_bidi={0}))
    assert blocks[0]["direction"] == "rtl"


def test_missing_word_bidi_uses_unicode_fallback():
    blocks = _blocks(_docx_bytes(["English text"], add_paragraph_properties=True))
    assert blocks[0]["direction"] == "ltr"


def test_empty_docx_has_specific_error():
    with pytest.raises(TextExtractError, match="empty_document") as error:
        extract_pages_geometry(_docx_bytes([]), "empty.docx")
    assert error.value.code == "empty_document"


def test_docx_without_meaningful_text_has_specific_error():
    with pytest.raises(TextExtractError, match="no_meaningful_text") as error:
        extract_pages_geometry(_docx_bytes(["   "]), "blank.docx")
    assert error.value.code == "no_meaningful_text"


def test_corrupted_docx_bytes_have_specific_error():
    with pytest.raises(TextExtractError, match="corrupted") as error:
        extract_pages_geometry(b"not a zip package", "broken.docx")
    assert error.value.code == "corrupted"


def test_encrypted_or_legacy_office_container_is_unsupported():
    ole_header = bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 32
    with pytest.raises(TextExtractError, match="encrypted_or_unsupported_docx") as error:
        extract_pages_geometry(ole_header, "encrypted.docx")
    assert error.value.code == "encrypted_or_unsupported_docx"


def test_unexpected_docx_failure_is_logged_without_document_data(caplog):
    with patch("app.services.textextract._docx_blocks", side_effect=AttributeError("internal failure")):
        with pytest.raises(TextExtractError, match="extraction_failed") as error:
            extract_pages_geometry(b"sensitive document contents", "private-person.docx")

    assert error.value.code == "extraction_failed"
    assert "AttributeError" in caplog.text
    assert "sensitive document contents" not in caplog.text
    assert "private-person.docx" not in caplog.text


def test_paragraph_order_is_preserved():
    expected = ["First paragraph", "الفقرة الثانية", "Third paragraph"]
    blocks = _blocks(_docx_bytes(expected))
    assert [block["text"] for block in blocks] == expected


@pytest.mark.parametrize("filename", ["valid.docx", "عقد-صحيح.docx"])
def test_http_upload_accepts_valid_docx_and_creates_initial_version(filename):
    db = MagicMock()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    try:
        with (
            patch("app.routers.contracts.storage.save", return_value=f"stored-{filename}"),
            patch("app.services.versions.create_initial_version") as create_initial_version,
        ):
            response = TestClient(app).post(
                "/api/contracts",
                files={
                    "file": (
                        filename,
                        _docx_bytes(["Valid contract paragraph"], add_paragraph_properties=True),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                headers={"Authorization": "Bearer demo-secret-token"},
            )
    finally:
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 201
    contract = db.add.call_args.args[0]
    assert contract.file_url == f"stored-{filename}"
    assert db.commit.called
    create_initial_version.assert_called_once_with(contract, db, actor="upload")
